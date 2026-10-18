import asyncio
import os
import io
import sys
import openpyxl
from playwright.async_api import async_playwright

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

async def run_tests():
    async with async_playwright() as p:
        browser = await p.chromium.launch(channel="msedge", headless=True)
        context = await browser.new_context()
        page = await context.new_page()

        print("[1] Logging in as teacher...")
        res = await page.request.post(
            "http://127.0.0.1:8000/api/teacher/login",
            data={"username": "teacher", "password": "bagalau2026"}
        )
        assert res.ok, f"Login failed: {res.status}"
        login_data = await res.json()
        token = login_data["token"]
        print(f"Logged in successfully! Teacher: {login_data['teacher']['full_name']}")

        await context.add_cookies([{
            "name": "bagalau_teacher_session",
            "value": token,
            "domain": "127.0.0.1",
            "path": "/"
        }])

        print("[2] Opening /teacher page...")
        await page.goto("http://127.0.0.1:8000/teacher")
        await page.wait_for_selector("#teacher-fullname")

        print("[3] Switching to launches tab...")
        await page.click('button[data-tab="launches"]')
        await page.wait_for_timeout(1000)

        # Check launch rows in table
        rows = await page.query_selector_all("#launches-table-body tr")
        print(f"Found {len(rows)} launch rows")
        assert len(rows) > 0, "No launches found in table"

        monitor_btn = await rows[0].query_selector('button:has-text("Мониторинг")')
        assert monitor_btn, "No monitor button found in first launch row"
        await monitor_btn.click()
        await page.wait_for_timeout(1000)
        print("Opened launch live view!")

        print("[4] In live monitor, switching to 'Результаты и проверка'...")
        await page.click('button[data-subtab="tab-live-results"]')
        await page.wait_for_timeout(1000)

        # Verify filters exist
        class_sel = await page.query_selector("#filter-results-class")
        letter_sel = await page.query_selector("#filter-results-letter")
        quarter_sel = await page.query_selector("#filter-results-quarter")
        var_sel = await page.query_selector("#filter-results-variant")
        search_in = await page.query_selector("#filter-results-search")
        reset_btn = await page.query_selector('button:has-text("Сбросить")')
        sort_sel = await page.query_selector("#sort-results-by")

        assert class_sel, "Missing #filter-results-class"
        assert letter_sel, "Missing #filter-results-letter"
        assert quarter_sel, "Missing #filter-results-quarter"
        assert var_sel, "Missing #filter-results-variant"
        assert search_in, "Missing #filter-results-search"
        assert reset_btn, "Missing reset button"
        assert not sort_sel, "#sort-results-by should NOT exist"
        print("[OK] Filter bar verification passed!")

        # Check students table
        student_rows = await page.query_selector_all("#grading-table-container tbody tr")
        print(f"Grading table contains {len(student_rows)} student rows")

        if len(student_rows) > 0:
            view_btn = await student_rows[0].query_selector('button:has-text("Просмотр")')
            print_btn = await student_rows[0].query_selector('button:has-text("Печать в PDF")')
            assert view_btn, "Missing 'Просмотр' button"
            assert print_btn, "Missing 'Печать в PDF' button"

            print("[5] Opening student grading modal...")
            await view_btn.click()
            await page.wait_for_timeout(1000)

            modal = await page.query_selector("#student-grading-modal.open")
            assert modal, "Modal did not open!"

            modal_text = await page.text_content("#student-grading-modal")
            assert "Код подтверждения" not in modal_text, "Confirmation code MUST NOT be in modal!"
            assert "ИТОГОВЫЙ БАЛЛ" in modal_text, "ИТОГОВЫЙ БАЛЛ missing in modal"
            assert "Время:" in modal_text or "Потраченное время" in modal_text, "Duration missing in modal"
            print("[OK] Student grading modal verified: confirmation code hidden, stats displayed!")

            # Check tasks inside modal
            task_cards = await page.query_selector_all("#grade-modal-tasks .print-task-card")
            print(f"Modal displays {len(task_cards)} task cards")

            # Check that options / answers are rendered cleanly
            card_html = await task_cards[0].inner_html() if task_cards else ""
            assert "Задание" in card_html
            print("[OK] Task cards rendered properly inside modal!")

            # Take screenshot of modal
            os.makedirs("scratch", exist_ok=True)
            await page.screenshot(path="scratch/verified_modal.png")
            print("[OK] Saved modal screenshot to scratch/verified_modal.png")

            # Close modal
            await page.click("#close-grade-modal")
            await page.wait_for_timeout(500)

        print("[6] Testing student PDF print window generation...")
        eval_res = await page.evaluate("""
            () => {
                if (!gradingRawData || !gradingRawData.attempts || gradingRawData.attempts.length === 0) return { ok: false, reason: 'no attempts' };
                const att = gradingRawData.attempts[0];
                const classInfo = parseClassAndLetter(att.student_class);
                return {
                    ok: true,
                    name: att.student_name,
                    grade: classInfo.grade,
                    letter: classInfo.letter,
                    variant: att.variant
                };
            }
        """)
        print(f"Print sheet evaluation data: {eval_res}")
        assert eval_res.get("ok"), "Evaluate failed"

        print("[7] Testing Excel & CSV exports directly from server endpoint...")
        con_launches = await page.request.get("http://127.0.0.1:8000/api/launches")
        launches_list = await con_launches.json()
        assert len(launches_list) > 0
        lid = launches_list[0]["id"]

        xlsx_res = await page.request.get(f"http://127.0.0.1:8000/api/launches/{lid}/export.xlsx")
        assert xlsx_res.ok, f"Export xlsx failed: {xlsx_res.status}"
        xlsx_bytes = await xlsx_res.body()
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        print(f"[OK] Excel downloaded! Sheet count: {len(wb.sheetnames)}, Names: {wb.sheetnames}")
        assert "Все результаты" in wb.sheetnames
        ws = wb["Все результаты"]
        headers = [ws.cell(row=4, column=c).value for c in range(1, 8)]
        assert headers == ["№", "ФИО Ученика", "Класс", "Дата сдачи", "Итог", "Макс", "Процент"], f"Bad headers: {headers}"
        print("[OK] Excel Sheet 1 headers strictly match requirements!")

        csv_res = await page.request.get(f"http://127.0.0.1:8000/api/launches/{lid}/export.csv")
        assert csv_res.ok, f"Export csv failed: {csv_res.status}"
        csv_text = await csv_res.text()
        assert "ФИО Ученика" in csv_text
        assert "ДЕТАЛЬНЫЕ ОТВЕТЫ УЧЕНИКОВ" in csv_text
        print("[OK] CSV export format verified!")

        await browser.close()
        print("\n=== ALL CHECKS AND VERIFICATIONS PASSED SUCCESSFULLY! ===")

if __name__ == "__main__":
    asyncio.run(run_tests())
