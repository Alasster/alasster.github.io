import io
import csv
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TASK_TYPE_NAMES = {
    "single_choice": "Один правильный выбор",
    "multi_choice": "Множественный выбор",
    "dropdown": "Выпадающий список",
    "true_false": "Истина / Ложь",
    "short_text": "Краткий ответ",
    "fill_blank": "Вставка пропусков",
    "number": "Числовой ответ",
    "matching": "Установление соответствия",
    "group_classification": "Распределение по категориям",
    "ordering": "Упорядочивание",
    "text_manual": "Развёрнутый ответ"
}

def format_answer_value(val) -> str:
    if val is None or val == "":
        return "Нет ответа"
    if isinstance(val, dict):
        return "\n".join(f"{k} → {v}" for k, v in val.items())
    if isinstance(val, list):
        return ", ".join(str(x) for x in val)
    return str(val)

def format_submitted_date(dt_str: str) -> str:
    if not dt_str or dt_str == "-":
        return "—"
    try:
        from datetime import datetime
        dt = datetime.fromisoformat(dt_str)
        return dt.strftime("%d.%m.%Y %H:%M:%S")
    except Exception:
        return str(dt_str)

def get_attempt_total_score(att: dict, answers_by_attempt: dict) -> float:
    att_answers = answers_by_attempt.get(att.get("id"), {})
    if att_answers:
        return round(sum(float(a.get("final_score") if a and a.get("final_score") is not None else 0.0) for a in att_answers.values()), 1)
    if att.get("total_score") is not None:
        try:
            return round(float(att["total_score"]), 1)
        except Exception:
            pass
    return 0.0

def generate_csv_report(launch_info: dict, attempts: list, tasks: list, answers_by_attempt: dict) -> str:
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")

    # Заголовок
    writer.writerow(["Отчёт по результатам проведения"])
    writer.writerow(["Работа:", launch_info.get("title", "")])
    writer.writerow(["Предмет:", launch_info.get("subject", "")])
    writer.writerow(["Класс:", launch_info.get("grade_level", "")])
    writer.writerow([])

    # Шапка таблицы результатов
    headers = ["№", "ФИО Ученика", "Класс", "Дата сдачи", "Итог", "Макс", "Процент"]
    writer.writerow(headers)

    sorted_attempts = sorted(attempts, key=lambda a: (a.get("student_name") or "").lower())

    for idx, att in enumerate(sorted_attempts, 1):
        stu_var = str(att.get("variant") or "1")
        stu_tasks = [t for t in tasks if str(t.get("variant") or "1") in (stu_var, "all", "1,2")] or tasks
        max_possible = sum(float(t.get("max_score", 1)) for t in stu_tasks)
        total_score = get_attempt_total_score(att, answers_by_attempt)
        percent = round((total_score / max_possible * 100), 1) if max_possible > 0 else 0

        writer.writerow([
            idx,
            att.get("student_name", ""),
            att.get("student_class", ""),
            format_submitted_date(att.get("submitted_at")),
            total_score,
            max_possible,
            f"{percent}%"
        ])

    writer.writerow([])
    writer.writerow(["--- ДЕТАЛЬНЫЕ ОТВЕТЫ УЧЕНИКОВ ПО ЗАДАНИЯМ ---"])

    for idx, att in enumerate(sorted_attempts, 1):
        stu_var = str(att.get("variant") or "1")
        stu_tasks = [t for t in tasks if str(t.get("variant") or "1") in (stu_var, "all", "1,2")] or tasks
        max_possible = sum(float(t.get("max_score", 1)) for t in stu_tasks)
        total_score = get_attempt_total_score(att, answers_by_attempt)
        percent = round((total_score / max_possible * 100), 1) if max_possible > 0 else 0
        att_answers = answers_by_attempt.get(att["id"], {})

        writer.writerow([])
        writer.writerow([f"Ученик {idx}:", att.get("student_name", ""), "Класс:", att.get("student_class", ""), "Вариант:", stu_var, "Итог:", f"{total_score} из {max_possible} ({percent}%)"])
        writer.writerow(["№ задания", "Тип задания", "Вопрос", "Ответ ученика", "Правильный ответ", "Балл", "Макс", "Статус"])

        for t_order, t in enumerate(stu_tasks, 1):
            ans = att_answers.get(t["id"], {})
            ans_parsed = ans.get("answer_parsed") if ans else None
            score = float(ans.get("final_score") if ans and ans.get("final_score") is not None else 0.0)
            t_max = float(t.get("max_score", 1))
            status_text = "Верно" if score >= t_max else ("Частично" if score > 0 else "Неверно")

            writer.writerow([
                t.get("order_index", t_order),
                TASK_TYPE_NAMES.get(t.get("question_type"), t.get("question_type", "")),
                t.get("question_text", ""),
                format_answer_value(ans_parsed).replace("\n", " | "),
                format_answer_value(t.get("answer_key")).replace("\n", " | "),
                score,
                t_max,
                status_text
            ])

    return output.getvalue()

def generate_excel_report(launch_info: dict, attempts: list, tasks: list, answers_by_attempt: dict) -> bytes:
    wb = openpyxl.Workbook()

    # Стили
    title_font = Font(name="Calibri", size=14, bold=True, color="1F2937")
    subtitle_font = Font(name="Calibri", size=11, bold=True, color="4B5563")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    accent_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    green_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    yellow_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    wrap_left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # -------------------------------------------------------------
    # ЛИСТ 1: "Все результаты"
    # Колонки: Номер, ФИО, Класс, Дата сдачи, Итог, Макс, Процент
    # -------------------------------------------------------------
    ws_all = wb.active
    ws_all.title = "Все результаты"

    ws_all["A1"] = f"Ведомость результатов: {launch_info.get('title', 'Работа')}"
    ws_all["A1"].font = title_font
    ws_all["A2"] = f"Предмет: {launch_info.get('subject', '')} | Класс: {launch_info.get('grade_level', '')} | Тип: {launch_info.get('exam_type', 'СОР')}"
    ws_all["A2"].font = subtitle_font

    summary_headers = ["№", "ФИО Ученика", "Класс", "Дата сдачи", "Итог", "Макс", "Процент"]
    for c_idx, h in enumerate(summary_headers, 1):
        cell = ws_all.cell(row=4, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    sorted_attempts = sorted(attempts, key=lambda a: (a.get("student_name") or "").lower())

    curr_row = 5
    for idx, att in enumerate(sorted_attempts, 1):
        stu_var = str(att.get("variant") or "1")
        stu_tasks = [t for t in tasks if str(t.get("variant") or "1") in (stu_var, "all", "1,2")] or tasks
        max_possible = sum(float(t.get("max_score", 1)) for t in stu_tasks)
        total_score = get_attempt_total_score(att, answers_by_attempt)
        percent = round((total_score / max_possible * 100), 1) if max_possible > 0 else 0

        ws_all.cell(row=curr_row, column=1, value=idx).alignment = center_align
        ws_all.cell(row=curr_row, column=2, value=att.get("student_name", "")).alignment = left_align
        ws_all.cell(row=curr_row, column=3, value=att.get("student_class", "")).alignment = center_align
        ws_all.cell(row=curr_row, column=4, value=format_submitted_date(att.get("submitted_at"))).alignment = center_align

        c_tot = ws_all.cell(row=curr_row, column=5, value=total_score)
        c_tot.alignment = center_align
        c_tot.font = Font(bold=True)
        c_tot.fill = accent_fill

        c_max = ws_all.cell(row=curr_row, column=6, value=max_possible)
        c_max.alignment = center_align

        c_pct = ws_all.cell(row=curr_row, column=7, value=f"{percent}%")
        c_pct.alignment = center_align
        c_pct.font = Font(bold=True)

        for c_idx in range(1, 8):
            ws_all.cell(row=curr_row, column=c_idx).border = thin_border

        curr_row += 1

    # -------------------------------------------------------------
    # СЛЕДУЮЩИЕ ЛИСТЫ: Отдельно для каждого ученика
    # -------------------------------------------------------------
    used_sheet_titles = set(["все результаты"])
    for idx, att in enumerate(sorted_attempts, 1):
        clean_name = re.sub(r'[\\/*?:\[\]]', '', att.get("student_name") or f"Ученик {idx}").strip()
        base_title = f"{idx}. {clean_name}"[:28]
        sheet_title = base_title
        dup_count = 1
        while sheet_title.lower() in used_sheet_titles:
            sheet_title = f"{base_title[:25]}_{dup_count}"
            dup_count += 1
        used_sheet_titles.add(sheet_title.lower())

        ws_stu = wb.create_sheet(title=sheet_title)

        stu_var = str(att.get("variant") or "1")
        stu_tasks = [t for t in tasks if str(t.get("variant") or "1") in (stu_var, "all", "1,2")] or tasks
        max_possible = sum(float(t.get("max_score", 1)) for t in stu_tasks)
        total_score = get_attempt_total_score(att, answers_by_attempt)
        percent = round((total_score / max_possible * 100), 1) if max_possible > 0 else 0
        att_answers = answers_by_attempt.get(att["id"], {})

        # Шапка ученика
        ws_stu["A1"] = f"Работа ученика: {att.get('student_name', '')}"
        ws_stu["A1"].font = title_font
        ws_stu["A2"] = f"Класс: {att.get('student_class', '')} | Вариант: {stu_var} | Сдано: {format_submitted_date(att.get('submitted_at'))}"
        ws_stu["A2"].font = subtitle_font
        ws_stu["A3"] = f"Итоговый балл: {total_score} из {max_possible} ({percent}%)"
        ws_stu["A3"].font = Font(name="Calibri", size=12, bold=True, color="4F46E5")

        stu_headers = ["№", "Тип задания", "Вопрос", "Ответ ученика", "Правильный ответ", "Балл", "Макс", "Статус"]
        for c_idx, h in enumerate(stu_headers, 1):
            cell = ws_stu.cell(row=5, column=c_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        r_idx = 6
        for t_order, t in enumerate(stu_tasks, 1):
            ans = att_answers.get(t["id"], {})
            ans_parsed = ans.get("answer_parsed") if ans else None
            score = float(ans.get("final_score") if ans and ans.get("final_score") is not None else 0.0)
            t_max = float(t.get("max_score", 1))

            is_correct = score >= t_max
            is_partial = 0 < score < t_max
            status_text = "Верно" if is_correct else ("Частично" if is_partial else "Неверно")
            status_fill = green_fill if is_correct else (yellow_fill if is_partial else red_fill)

            ws_stu.cell(row=r_idx, column=1, value=t.get("order_index", t_order)).alignment = center_align
            ws_stu.cell(row=r_idx, column=2, value=TASK_TYPE_NAMES.get(t.get("question_type"), t.get("question_type", ""))).alignment = left_align
            ws_stu.cell(row=r_idx, column=3, value=t.get("question_text", "")).alignment = wrap_left_align
            ws_stu.cell(row=r_idx, column=4, value=format_answer_value(ans_parsed)).alignment = wrap_left_align
            ws_stu.cell(row=r_idx, column=5, value=format_answer_value(t.get("answer_key"))).alignment = wrap_left_align

            ws_stu.cell(row=r_idx, column=6, value=score).alignment = center_align
            ws_stu.cell(row=r_idx, column=7, value=t_max).alignment = center_align

            c_st = ws_stu.cell(row=r_idx, column=8, value=status_text)
            c_st.alignment = center_align
            c_st.font = Font(bold=True)
            c_st.fill = status_fill

            for col_idx in range(1, 9):
                ws_stu.cell(row=r_idx, column=col_idx).border = thin_border

            r_idx += 1

    # Авто-ширина колонок для всех листов
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                if "\n" in val:
                    val = max(val.split("\n"), key=len)
                max_len = max(max_len, len(val))
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
