"""Check answer contents in both export formats; API checks use a temporary DB.

Run: python -m unittest test_export_answers -v
"""
import csv
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import httpx
import openpyxl

import database
import server
from auth import create_teacher_session, destroy_teacher_session
from export import generate_csv_report, generate_excel_report, get_export_answer_value


class ExportAnswerTests(unittest.TestCase):
    def test_persisted_values_and_legacy_text(self):
        for value in ("Ответ ученика", ["А", "В"], {"Термин": "Описание"},
                      ["Первый", "Второй"], 0, False, None, ""):
            with self.subTest(value=value):
                self.assertEqual(get_export_answer_value({"answer_json": json.dumps(value)}), value)
        self.assertEqual(get_export_answer_value({"answer_json": "Старый текст без JSON"}), "Старый текст без JSON")
        self.assertEqual(get_export_answer_value({"answer_parsed": ["А", "Б"]}), ["А", "Б"])
        self.assertIsNone(get_export_answer_value({}))
        self.assertIsNone(get_export_answer_value({"answer_json": "null"}))
        self.assertEqual(get_export_answer_value({"answer_json": '"Сохранено"', "answer_parsed": "Устарело"}), "Сохранено")

    def test_both_reports_contain_answers_for_each_task_type(self):
        cases = [
            ("single_choice", "Принтер", "Принтер"),
            ("multi_choice", ["CPU", "RAM"], "CPU, RAM"),
            ("dropdown", "Python", "Python"),
            ("true_false", "Верно", "Верно"),
            ("short_text", "Собственный ответ", "Собственный ответ"),
            ("fill_blank", ["один", "два"], "один, два"),
            ("number", 0, "0"),
            ("matching", {"A": "1", "Б": "2"}, "A → 1\nБ → 2"),
            ("group_classification", {"Ввод": ["Мышь", "Клавиатура"]}, "Ввод → ['Мышь', 'Клавиатура']"),
            ("ordering", ["Шаг 3", "Шаг 1", "Шаг 2"], "Шаг 3, Шаг 1, Шаг 2"),
            ("text_manual", "Развёрнутый ответ; с разделителем\nи новой строкой", "Развёрнутый ответ; с разделителем\nи новой строкой"),
            ("number", False, "False"),
            ("short_text", None, "Нет ответа"),
        ]
        tasks = [{"id": str(i), "question_type": kind, "question_text": f"Вопрос {i}",
                  "max_score": 1, "variant": "1", "answer_key": "Ключ"}
                 for i, (kind, _, _) in enumerate(cases)]
        answers = {"a": {str(i): {"answer_json": json.dumps(value, ensure_ascii=False), "final_score": 1}
                         for i, (_, value, _) in enumerate(cases)}}
        attempts = [{"id": "a", "student_name": "Тестовый ученик", "student_class": "10 А", "variant": "1"}]
        workbook = openpyxl.load_workbook(io.BytesIO(generate_excel_report({}, attempts, tasks, answers)))
        rows = list(csv.reader(io.StringIO(generate_csv_report({}, attempts, tasks, answers)), delimiter=";"))
        csv_details = [row for row in rows if len(row) == 8 and row[2].startswith("Вопрос ")]
        self.assertEqual(len(csv_details), len(cases))
        for index, (_, _, expected) in enumerate(cases):
            with self.subTest(index=index):
                self.assertEqual(workbook.worksheets[1].cell(index + 6, 4).value, expected)
                self.assertEqual(csv_details[index][3], expected.replace("\n", " | "))
                self.assertEqual(workbook.worksheets[1].cell(index + 6, 6).value, 1)
        self.assertEqual(workbook.worksheets[0]["E5"].value, len(cases))
        workbook.close()


class ExportEndpointTests(unittest.IsolatedAsyncioTestCase):
    async def test_excel_and_csv_read_saved_answers_and_keep_variants_separate(self):
        temp = tempfile.TemporaryDirectory(prefix="bagalau-export-regression-")
        self.addCleanup(temp.cleanup)
        db_patch = patch.multiple(database, DB_DIR=temp.name,
                                  DB_PATH=str(Path(temp.name) / "bagalau.sqlite3"))
        db_patch.start()
        self.addCleanup(db_patch.stop)
        await database.init_db()
        tasks = [{"id": f"q{v}", "variant": v, "question_text": f"Задание варианта {v}",
                  "question_type": "short_text", "max_score": 2, "answer_key": "Ключ"}
                 for v in ("1", "2")]
        tasks.append({"id": "common", "variant": "all", "question_text": "Общее задание",
                      "question_type": "number", "max_score": 1, "answer_key": 0})
        now = "2026-10-19T10:00:00+00:00"
        async with database.get_db() as db:
            teacher = await (await db.execute("SELECT id FROM teachers LIMIT 1")).fetchone()
            teacher_id = teacher["id"]
            await db.execute("""INSERT INTO tests
                (id, teacher_id, title, exam_type, subject, grade_level, quarter, created_at, updated_at)
                VALUES ('test', ?, 'Экспорт ответов', 'СОЧ', 'Информатика', '10', 1, ?, ?)""",
                (teacher_id, now, now))
            await db.execute("""INSERT INTO launches
                (id, test_id, teacher_id, title, status, allowed_classes_json, snapshot_tasks_json, created_at)
                VALUES ('launch', 'test', ?, 'Экспорт ответов', 'active', '[]', ?, ?)""",
                (teacher_id, json.dumps(tasks, ensure_ascii=False), now))
            for v in ("1", "2"):
                await db.execute("""INSERT INTO student_attempts
                    (id, launch_id, student_name, student_class, variant, attempt_token, status, created_at)
                    VALUES (?, 'launch', ?, '10 А', ?, ?, 'submitted', ?)""",
                    (f"a{v}", f"Ученик {v}", v, f"token{v}", now))
                for task_id, value, score in ((f"q{v}", f"Ответ ученика варианта {v}", 2), ("common", 0, 1)):
                    await db.execute("""INSERT INTO student_answers
                        (id, attempt_id, task_id, answer_json, final_score, updated_at)
                        VALUES (?, ?, ?, ?, ?, ?)""",
                        (f"a{v}-{task_id}", f"a{v}", task_id, json.dumps(value, ensure_ascii=False), score, now))
            await db.commit()
        token = create_teacher_session(teacher_id)
        self.addCleanup(destroy_teacher_session, token)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app), base_url="http://test") as client:
            headers = {"Authorization": "Bearer " + token}
            xlsx = await client.get("/api/launches/launch/export.xlsx", headers=headers)
            csv_response = await client.get("/api/launches/launch/export.csv", headers=headers)
        self.assertEqual(xlsx.status_code, 200, xlsx.text if xlsx.status_code != 200 else "")
        self.assertEqual(csv_response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(xlsx.content))
        for index in (1, 2):
            sheet = workbook.worksheets[index]
            name_cell = workbook.worksheets[0].cell(index + 4, 2)
            self.assertEqual(name_cell.hyperlink.location, f"'{sheet.title}'!A1")
            self.assertIsNone(name_cell.hyperlink.target)
            self.assertEqual(sheet["D6"].value, f"Ответ ученика варианта {index}")
            self.assertEqual(sheet["D7"].value, "0")
            self.assertEqual(sheet.max_row, 7)
            self.assertEqual(workbook.worksheets[0].cell(index + 4, 5).value, 3)
        self.assertTrue(csv_response.content.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig")), delimiter=";"))
        details = [row for row in rows if len(row) == 8 and row[2].startswith(("Задание", "Общее"))]
        self.assertEqual([row[3] for row in details], ["Ответ ученика варианта 1", "0", "Ответ ученика варианта 2", "0"])
        workbook.close()


if __name__ == "__main__":
    unittest.main()
